from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx

NAVY = RGBColor(0x1F,0x3A,0x5F); TEAL=RGBColor(0x0F,0x76,0x6E); GREY=RGBColor(0x55,0x55,0x55)
doc = Document()
for s in doc.sections:
    s.left_margin=s.right_margin=Cm(2); s.top_margin=s.bottom_margin=Cm(1.8)

def set_lang(style):
    rpr = style.element.get_or_add_rPr()
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'),'en-CA'); lang.set(qn('w:eastAsia'),'en-CA'); rpr.append(lang)
st = doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5); set_lang(st)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
st.paragraph_format.space_after=Pt(4)
for n,sz,c in [('Heading 1',17,NAVY),('Heading 2',13.5,TEAL),('Heading 3',11.5,NAVY)]:
    h=doc.styles[n]; h.font.name='Calibri'; h.font.size=Pt(sz); h.font.color.rgb=c; h.font.bold=True; set_lang(h)
    h.element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); h.element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
# doc-level language default
settings = doc.settings.element

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexcolor); tcPr.append(sh)

def add_link(par, url, text=None, size=None):
    part = par.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), r_id)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    col = OxmlElement('w:color'); col.set(qn('w:val'),'0563C1'); rPr.append(col)
    u = OxmlElement('w:u'); u.set(qn('w:val'),'single'); rPr.append(u)
    if size:
        sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(int(size*2))); rPr.append(sz)
    r.append(rPr); t = OxmlElement('w:t'); t.text = text or url; t.set(qn('xml:space'),'preserve'); r.append(t); h.append(r)
    par._p.append(h)

def rich(par, text, size=None):
    """**bold** segments supported"""
    parts = text.split('**')
    for i,p in enumerate(parts):
        if not p: continue
        r = par.add_run(p); r.bold = (i%2==1)
        if size: r.font.size=Pt(size)
    return par

def P(text, style=None, size=None, italic=False, color=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    rich(p, text, size)
    if italic or color:
        for r in p.runs:
            r.italic = italic
            if color: r.font.color.rgb = color
    return p
def B(text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); rich(p,text); p.paragraph_format.space_after=Pt(2); return p
def N(text):
    p = doc.add_paragraph(style='List Number'); rich(p,text); p.paragraph_format.space_after=Pt(2); return p
H1=lambda t: doc.add_heading(t,1); H2=lambda t: doc.add_heading(t,2); H3=lambda t: doc.add_heading(t,3)

def table(headers, rows, widths, fs=8.5, header_fill='1F3A5F'):
    t = doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(fs+0.5); r.font.color.rgb=RGBColor(255,255,255); shade(c,header_fill)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; c.text=''; p=c.paragraphs[0]
            if isinstance(val, tuple):  # (name, url, extra)
                name,url,extra = val
                r=p.add_run(name); r.bold=True; r.font.size=Pt(fs)
                p2=c.add_paragraph(); add_link(p2,url,url.replace('https://',''),size=fs-0.5)
                if extra:
                    p3=c.add_paragraph(); r3=p3.add_run(extra); r3.font.size=Pt(fs-0.5); r3.italic=True
            else:
                rich(p,val,fs)
            if ri%2==1: shade(c,'EEF3F8')
    t.autofit=False
    for row in t.rows:
        for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    for i,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[i]*567)))
    doc.add_paragraph()
    return t

def kv_table(pairs, fs=9.5):
    t=doc.add_table(rows=0, cols=2); t.style='Table Grid'
    for k,v in pairs:
        cells=t.add_row().cells
        cells[0].text=''; r=cells[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(fs); shade(cells[0],'E6F2F1')
        cells[1].text=''; rich(cells[1].paragraphs[0], v, fs)
    for row in t.rows:
        row.cells[0].width=Cm(4.2); row.cells[1].width=Cm(12.8)
    t.autofit=False
    for i,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int([4.2,12.8][i]*567)))
    doc.add_paragraph()

